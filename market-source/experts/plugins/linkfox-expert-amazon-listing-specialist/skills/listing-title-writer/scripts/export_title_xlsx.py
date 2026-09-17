#!/usr/bin/env python3
"""
listing-title-writer Excel 导出器 — 把 Workbench JSON 变成一张能直接看的表。

**列定义不在本文件维护**：来自 payload 里的 `columns`（由 save_title_output.py 的
`WORKBENCH_COLUMNS` 生成），前端 TitleWorkbenchRenderer 用的是同一份。
改列只改 `save_title_output.py`，Excel 与工作台同步生效。

一行一个商品，一张 sheet：
  行号 / SKU / ASIN / 原标题 / 新标题 / 标题字符数 / Item Highlights /
  Highlights 字符数 / 埋词说明 / 改写逻辑 / 状态 / 校验说明

Usage:
  python3 scripts/export_title_xlsx.py <workbench.json> [--source-file 原表.xlsx|csv|tsv]
      [--out /abs/path/out.xlsx]

输出 stdout：`XLSX artifact: <绝对路径>`
"""

from __future__ import annotations

import argparse
import csv
import json
import os
import sys
import time
from pathlib import Path
from typing import Any

STATUS_LABEL = {"ok": "已通过", "review": "待复核", "failed": "未通过"}

# payload 缺少 columns 时的兜底（形状与 save_title_output.WORKBENCH_COLUMNS 一致）
FALLBACK_COLUMNS: list[dict[str, Any]] = [
    {"key": "row_index", "title": "行号", "excelWidth": 6},
    {"key": "sku", "title": "SKU", "excelWidth": 14},
    {"key": "asin", "title": "ASIN", "excelWidth": 13},
    {"key": "legacy_title", "title": "原标题", "excelWidth": 40},
    {"key": "title", "title": "新标题", "excelWidth": 40},
    {"key": "title_char_count", "title": "标题字符数", "excelWidth": 11, "render": "charMeter", "maxFrom": "row.policy.title_max"},
    {"key": "item_highlights", "title": "Item Highlights", "excelWidth": 40},
    {"key": "item_highlights_char_count", "title": "Highlights 字符数", "excelWidth": 13, "render": "charMeter", "maxFrom": "row.policy.item_highlights_max"},
    {"key": "keyword_notes", "title": "埋词说明", "excelWidth": 38, "render": "multiline"},
    {"key": "rationale", "title": "改写逻辑", "excelWidth": 34, "render": "multiline"},
    {"key": "status", "title": "状态", "excelWidth": 9, "render": "statusTag"},
    {"key": "validation_notes", "title": "校验说明", "excelWidth": 30, "render": "multiline"},
]


# --- 数据提取 ------------------------------------------------------------


def _dict(value: Any) -> dict[str, Any]:
    return value if isinstance(value, dict) else {}


def _list(value: Any) -> list[Any]:
    return value if isinstance(value, list) else []


def load_rows(payload: Any) -> list[dict[str, Any]]:
    """接受 Workbench 信封、裸 rows[]、或单条 row。"""
    if isinstance(payload, list):
        rows = payload
    elif isinstance(payload, dict):
        rows = payload.get("rows") or payload.get("items") or payload.get("results")
        if not isinstance(rows, list):
            rows = [payload]
    else:
        rows = []
    return [r for r in rows if isinstance(r, dict)]


def load_columns(payload: Any) -> list[dict[str, Any]]:
    columns = _dict(payload).get("columns") if isinstance(payload, dict) else None
    picked = [c for c in _list(columns) if isinstance(c, dict) and c.get("key")]
    return picked or FALLBACK_COLUMNS


def _resolve_max(row: dict[str, Any], max_from: str | None) -> int | None:
    """解析 columns 里的 `row.policy.title_max` 这类引用。"""
    if not max_from:
        return None
    node: Any = row
    for part in max_from.split(".")[1:]:  # 跳过开头的 "row"
        node = _dict(node).get(part)
    return node if isinstance(node, int) and not isinstance(node, bool) else None


def cell_value(row: dict[str, Any], column: dict[str, Any]) -> Any:
    """按共享列契约取值；rows[] 的扁平镜像已由保存脚本备好。"""
    key = str(column.get("key"))
    value = row.get(key)

    if key == "status":
        return STATUS_LABEL.get(str(value or "ok"), str(value or "ok"))
    if column.get("render") == "charMeter":
        count = value if isinstance(value, int) else 0
        limit = _resolve_max(row, column.get("maxFrom"))
        return f"{count} / {limit}" if limit else count
    if isinstance(value, list):
        return "\n".join(str(item) for item in value if str(item).strip())
    if isinstance(value, dict):
        return json.dumps(value, ensure_ascii=False)
    return value


def is_over_limit(row: dict[str, Any], column: dict[str, Any]) -> bool:
    if column.get("render") != "charMeter":
        return False
    limit = _resolve_max(row, column.get("maxFrom"))
    count = row.get(str(column.get("key")))
    return bool(limit and isinstance(count, int) and count > limit)


# --- Excel 写出 ----------------------------------------------------------


def _styles():
    from openpyxl.styles import Alignment, Font, PatternFill

    return {
        "header_font": Font(bold=True, color="FFFFFF"),
        "header_fill": PatternFill("solid", fgColor="1F4E79"),
        # 追加到用户原表右侧的新增列用另一种颜色区分：深绿表头 + 浅绿单元格
        "new_header_fill": PatternFill("solid", fgColor="2E7D32"),
        "new_cell_fill": PatternFill("solid", fgColor="E8F5E9"),
        "wrap": Alignment(vertical="top", wrap_text=True),
        "over_font": Font(color="C00000", bold=True),
        "failed_fill": PatternFill("solid", fgColor="FCE4E4"),
        "review_fill": PatternFill("solid", fgColor="FFF6E0"),
    }


def _write_header(
    sheet,
    titles: list[str],
    widths: list[int],
    styles,
    *,
    new_from: int | None = None,
) -> None:
    """写表头；`new_from` 起（1-based 列号）视为本次新增列，用另一种颜色标注。"""
    from openpyxl.utils import get_column_letter

    for idx, title in enumerate(titles, start=1):
        cell = sheet.cell(row=1, column=idx, value=title)
        cell.font = styles["header_font"]
        cell.fill = (
            styles["new_header_fill"]
            if new_from is not None and idx >= new_from
            else styles["header_fill"]
        )
        cell.alignment = styles["wrap"]
        sheet.column_dimensions[get_column_letter(idx)].width = widths[idx - 1]
    sheet.freeze_panes = "A2"


def _write_appended_header(sheet, columns: list[dict[str, Any]], offset: int, styles) -> None:
    """只格式化新增列，避免覆盖用户原工作簿的表头样式和列宽。"""
    from openpyxl.utils import get_column_letter

    for position, column in enumerate(columns, start=offset + 1):
        cell = sheet.cell(row=1, column=position, value=str(column.get("title") or ""))
        cell.font = styles["header_font"]
        cell.fill = styles["new_header_fill"]
        cell.alignment = styles["wrap"]
        sheet.column_dimensions[get_column_letter(position)].width = int(
            column.get("excelWidth") or 24
        )
    if sheet.freeze_panes is None:
        sheet.freeze_panes = "A2"


def _status_fill(row: dict[str, Any], styles):
    status = str(row.get("status") or "ok")
    if status == "failed":
        return styles["failed_fill"]
    if status == "review":
        return styles["review_fill"]
    return None


def _read_source_table(path: Path) -> tuple[list[str], list[list[Any]]]:
    """读原表，保留原始列，用于与生成结果横向拼接。"""
    suffix = path.suffix.lower()
    if suffix in {".csv", ".tsv", ".txt"}:
        delimiter = "\t" if suffix == ".tsv" else ","
        with open(path, encoding="utf-8-sig", newline="") as f:
            table = list(csv.reader(f, delimiter=delimiter))
        if not table:
            return [], []
        return [str(h) for h in table[0]], [list(r) for r in table[1:]]
    if suffix == ".xlsx":
        import openpyxl  # type: ignore

        workbook = openpyxl.load_workbook(path, data_only=True)
        sheet = workbook.active
        table = [list(row) for row in sheet.iter_rows(values_only=True)]
        if not table:
            return [], []
        return [str(h) if h is not None else "" for h in table[0]], [list(r) for r in table[1:]]
    raise RuntimeError(f"不支持的源表格式：{suffix}（.xls 请先另存为 .xlsx）")


def _already_in_source(
    column: dict[str, Any],
    rows: list[dict[str, Any]],
    src_rows: list[list[Any]],
) -> bool:
    """原表已经有同样内容的列时不重复追加（典型：用户表自带 Original Title）。

    只对原标题做判定——它是唯一会与用户上传表重复的列；靠值比对而不是列名，
    因为用户的列名可能是 Original Title / 旧标题 / Title 等任意写法。
    """
    if column.get("key") != "legacy_title":
        return False

    source_cells = {
        str(cell).strip()
        for src_row in src_rows
        for cell in src_row
        if isinstance(cell, str) and cell.strip()
    }
    if not source_cells:
        return False

    values = [str(row.get("legacy_title") or "").strip() for row in rows]
    values = [v for v in values if v]
    return bool(values) and all(value in source_cells for value in values)


def _target_row(row: dict[str, Any]) -> int:
    """写回原表的目标行号：优先用批量保存脚本解析好的 `target_row`。"""
    if row.get("writeback_blocked") is True:
        return 0
    for key in ("target_row", "source_row"):
        value = row.get(key)
        if isinstance(value, int) and not isinstance(value, bool) and value >= 2:
            return value
    row_index = row.get("row_index")
    if isinstance(row_index, int) and not isinstance(row_index, bool):
        return row_index + 1 if row_index >= 1 else 0
    return 0


def build_workbook(payload: Any, source_file: str | None = None):
    import openpyxl  # type: ignore

    rows = load_rows(payload)
    columns = load_columns(payload)
    styles = _styles()

    if source_file:
        source_path = Path(source_file).expanduser().resolve()
        preserve_workbook = source_path.suffix.lower() == ".xlsx"
        if preserve_workbook:
            # 在源工作簿副本上追加：保留公式、样式、合并单元格和其它 sheet。
            workbook = openpyxl.load_workbook(source_path)
            sheet = workbook.active
            src_headers = [
                str(cell.value) if cell.value is not None else ""
                for cell in sheet[1]
            ]
            src_rows = [list(row) for row in sheet.iter_rows(min_row=2, values_only=True)]
        else:
            src_headers, src_rows = _read_source_table(source_path)
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.title = "标题结果"
        appended = [
            c
            for c in columns
            if str(c.get("title")) not in src_headers and c.get("key") != "row_index"
        ]
        appended = [c for c in appended if not _already_in_source(c, rows, src_rows)]
        titles = [h or f"列{i + 1}" for i, h in enumerate(src_headers)] + [str(c.get("title")) for c in appended]
        if preserve_workbook:
            _write_appended_header(sheet, appended, len(src_headers), styles)
        else:
            widths = [18] * len(src_headers) + [int(c.get("excelWidth") or 24) for c in appended]
            _write_header(sheet, titles, widths, styles, new_from=len(src_headers) + 1)
            for r, src_row in enumerate(src_rows, start=2):
                for c, value in enumerate(src_row, start=1):
                    sheet.cell(row=r, column=c, value=value).alignment = styles["wrap"]

        offset = len(src_headers)
        status_index = next((i for i, c in enumerate(appended) if c.get("key") == "status"), None)
        skipped: list[int] = []
        for row in rows:
            excel_row = _target_row(row)
            if excel_row < 2 or excel_row > len(src_rows) + 1:
                skipped.append(int(row.get("row_index") or 0))
                continue
            _write_row(sheet, excel_row, offset, appended, row, styles, highlight=True)
            fill = _status_fill(row, styles)
            if fill and status_index is not None:
                sheet.cell(row=excel_row, column=offset + status_index + 1).fill = fill
        if skipped:
            print(f"⚠️ {len(skipped)} 行的 row_index 超出原表范围，已跳过写回：{skipped[:10]}", file=sys.stderr)
        sheet.auto_filter.ref = (
            f"A1:{sheet.cell(row=1, column=len(titles)).coordinate}{len(src_rows) + 1}"
        )
        return workbook

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "标题结果"
    titles = [str(c.get("title")) for c in columns]
    widths = [int(c.get("excelWidth") or 24) for c in columns]
    _write_header(sheet, titles, widths, styles)
    status_index = next((i for i, c in enumerate(columns) if c.get("key") == "status"), None)
    for r, row in enumerate(rows, start=2):
        _write_row(sheet, r, 0, columns, row, styles)
        fill = _status_fill(row, styles)
        if fill and status_index is not None:
            sheet.cell(row=r, column=status_index + 1).fill = fill
    if rows:
        sheet.auto_filter.ref = f"A1:{sheet.cell(row=1, column=len(titles)).coordinate}{len(rows) + 1}"
    return workbook


def _write_row(
    sheet,
    excel_row: int,
    offset: int,
    columns: list[dict[str, Any]],
    row: dict[str, Any],
    styles,
    *,
    highlight: bool = False,
) -> None:
    for idx, column in enumerate(columns, start=offset + 1):
        cell = sheet.cell(row=excel_row, column=idx, value=cell_value(row, column))
        cell.alignment = styles["wrap"]
        if highlight:
            cell.fill = styles["new_cell_fill"]
        if is_over_limit(row, column):
            cell.font = styles["over_font"]


def resolve_output_path(explicit: str | None) -> Path:
    if explicit:
        path = Path(explicit).expanduser().resolve()
        path.parent.mkdir(parents=True, exist_ok=True)
        return path
    script_dir = os.path.dirname(os.path.abspath(__file__))
    sys.path.insert(0, script_dir)
    try:
        from linkfox_paths import resolve_report_path  # type: ignore

        return Path(resolve_report_path("listing-title-workbench", time.time(), "xlsx")).resolve()
    except Exception:  # noqa: BLE001 — 本地调试时回退到当前目录
        return Path(os.getcwd(), "listing-title-workbench.xlsx").resolve()


def export(payload: Any, *, out: str | None = None, source_file: str | None = None) -> Path:
    workbook = build_workbook(payload, source_file=source_file)
    output = resolve_output_path(out)
    workbook.save(output)
    return output


def main() -> None:
    parser = argparse.ArgumentParser(description="把 listing-title-writer 结果导出为可批量查看的 Excel")
    parser.add_argument("json_file", help="Workbench JSON（save_title_*_output.py 落盘结果）")
    parser.add_argument("--source-file", help="可选原表（.xlsx/.csv/.tsv）：保留原始列并在右侧追加结果")
    parser.add_argument("--out", help="可选输出路径；默认写入会话 reports 目录")
    args = parser.parse_args()

    with open(args.json_file, encoding="utf-8") as f:
        payload = json.load(f)
    rows = load_rows(payload)
    if not rows:
        print("没有可导出的行", file=sys.stderr)
        sys.exit(1)

    output = export(payload, out=args.out, source_file=args.source_file)
    statuses = [str(row.get("status") or "ok") for row in rows]
    print(f"XLSX artifact: {output}")
    print(
        f"  rows: {len(rows)} · ok: {statuses.count('ok')} · "
        f"review: {statuses.count('review')} · failed: {statuses.count('failed')}"
    )


if __name__ == "__main__":
    main()
