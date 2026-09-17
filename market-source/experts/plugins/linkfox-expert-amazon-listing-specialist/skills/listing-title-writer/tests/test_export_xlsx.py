from __future__ import annotations

import copy
import json
import sys
import tempfile
import unittest
from pathlib import Path


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from export_title_xlsx import (  # noqa: E402
    _target_row,
    build_workbook,
    cell_value,
    is_over_limit,
    load_columns,
    load_rows,
)
from save_title_output import (  # noqa: E402
    WORKBENCH_COLUMNS,
    _validate,
    attach_xlsx_artifact,
    build_workbench_payload,
)

BATCH = json.loads((ROOT / "tests" / "fixtures" / "title-batch-sample.json").read_text("utf-8"))
SOURCE_CSV = ROOT / "tests" / "fixtures" / "source-table-sample.csv"

EXPECTED_TITLES = [
    "行号",
    "SKU",
    "ASIN",
    "原标题",
    "新标题",
    "标题字符数",
    "Item Highlights",
    "Highlights 字符数",
    "埋词说明",
    "改写逻辑",
    "状态",
    "校验说明",
]


def _workbench(payload: dict) -> dict:
    validated = []
    for position, row in enumerate(copy.deepcopy(payload)["rows"], start=1):
        row.setdefault("row_index", position)
        errors, warnings = _validate(row)
        validated.append((row, errors, warnings))
    return build_workbench_payload(validated)


class SharedColumnContractTest(unittest.TestCase):
    def test_columns_come_from_the_payload_not_the_exporter(self) -> None:
        payload = _workbench(BATCH)
        self.assertEqual([c["title"] for c in payload["columns"]], EXPECTED_TITLES)
        self.assertEqual(load_columns(payload), payload["columns"])
        self.assertEqual([c["title"] for c in WORKBENCH_COLUMNS], EXPECTED_TITLES)

    def test_fallback_columns_used_when_payload_has_none(self) -> None:
        columns = load_columns({"rows": []})
        self.assertEqual([c["title"] for c in columns], EXPECTED_TITLES)

    def test_load_rows_accepts_workbench_bare_rows_and_single_row(self) -> None:
        rows = _workbench(BATCH)["rows"]
        self.assertEqual(len(load_rows({"rows": rows})), 3)
        self.assertEqual(len(load_rows(rows)), 3)
        self.assertEqual(len(load_rows(rows[0])), 1)


class KeywordNotesTest(unittest.TestCase):
    def setUp(self) -> None:
        self.rows = _workbench(BATCH)["rows"]

    def test_notes_explain_placement_bucket_handoff_and_source(self) -> None:
        notes = self.rows[0]["keyword_notes"]
        self.assertIn("标题：dog water bottle（核心词）", notes)
        self.assertIn("Highlights：19 oz（属性词）、hiking（场景词）", notes)
        self.assertIn("交五点：leakproof dog water bottle", notes)
        self.assertIn("交后台：puppy water cup", notes)
        self.assertIn("词源：SIF 流量词", notes)

    def test_heuristic_source_is_labelled_as_no_traffic_data(self) -> None:
        self.assertIn("推导词（无流量数据）", self.rows[2]["keyword_notes"])

    def test_field_is_inferred_when_missing(self) -> None:
        payload = copy.deepcopy(BATCH)
        payload["rows"] = [payload["rows"][0]]
        payload["rows"][0]["used_keywords"] = [{"kw": "Leakproof"}, {"kw": "19oz"}]
        notes = _workbench(payload)["rows"][0]["keyword_notes"]
        self.assertIn("标题：Leakproof", notes)
        self.assertIn("Highlights：19oz", notes)


class CellValueTest(unittest.TestCase):
    def setUp(self) -> None:
        self.rows = _workbench(BATCH)["rows"]
        self.columns = {c["key"]: c for c in WORKBENCH_COLUMNS}

    def test_char_meter_shows_count_over_limit(self) -> None:
        self.assertEqual(cell_value(self.rows[0], self.columns["title_char_count"]), "43 / 75")
        self.assertEqual(cell_value(self.rows[2], self.columns["title_char_count"]), "94 / 75")

    def test_over_limit_detection(self) -> None:
        self.assertFalse(is_over_limit(self.rows[0], self.columns["title_char_count"]))
        self.assertTrue(is_over_limit(self.rows[2], self.columns["title_char_count"]))

    def test_status_is_localized_and_notes_are_joined(self) -> None:
        self.assertEqual(cell_value(self.rows[0], self.columns["status"]), "已通过")
        self.assertEqual(cell_value(self.rows[2], self.columns["status"]), "未通过")
        notes = cell_value(self.rows[2], self.columns["validation_notes"])
        self.assertIn("超长", notes)
        self.assertIn("\n", notes)


class ExportWorkbookTest(unittest.TestCase):
    def setUp(self) -> None:
        self.payload = _workbench(BATCH)

    def test_single_sheet_with_shared_columns(self) -> None:
        workbook = build_workbook(self.payload)
        self.assertEqual(workbook.sheetnames, ["标题结果"])

        sheet = workbook["标题结果"]
        self.assertEqual([cell.value for cell in sheet[1]], EXPECTED_TITLES)
        self.assertEqual(sheet.max_row, 4)
        self.assertEqual(sheet.freeze_panes, "A2")
        self.assertTrue(sheet.auto_filter.ref)

    def test_over_limit_char_count_is_marked_red(self) -> None:
        sheet = build_workbook(self.payload)["标题结果"]
        col = EXPECTED_TITLES.index("标题字符数") + 1
        normal = sheet.cell(row=2, column=col).font.color
        self.assertNotEqual(getattr(normal, "rgb", None), "00C00000")
        self.assertEqual(sheet.cell(row=4, column=col).font.color.rgb, "00C00000")

    def test_keyword_notes_reach_the_sheet(self) -> None:
        sheet = build_workbook(self.payload)["标题结果"]
        col = EXPECTED_TITLES.index("埋词说明") + 1
        self.assertIn("核心词", sheet.cell(row=2, column=col).value)

    def test_source_file_columns_are_preserved_and_appended(self) -> None:
        sheet = build_workbook(self.payload, source_file=str(SOURCE_CSV))["标题结果"]
        headers = [cell.value for cell in sheet[1]]

        self.assertEqual(headers[:4], ["SKU", "ASIN", "Legacy Title", "Category"])
        self.assertIn("新标题", headers)
        self.assertIn("埋词说明", headers)
        self.assertEqual(headers.count("SKU"), 1)
        self.assertEqual(sheet.cell(row=2, column=1).value, "PAWGO-19")
        self.assertEqual(
            sheet.cell(row=2, column=headers.index("新标题") + 1).value,
            "PawGo Leakproof Dog Water Bottle for Travel",
        )

    def test_blocked_writeback_never_falls_back_to_row_index(self) -> None:
        self.assertEqual(
            _target_row({"target_row": 2, "row_index": 1, "writeback_blocked": True}),
            0,
        )

    def test_xlsx_source_preserves_formula_style_and_other_sheets(self) -> None:
        import openpyxl
        from openpyxl.styles import PatternFill

        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "source.xlsx"
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.title = "原始数据"
            sheet.append(["SKU", "ASIN", "Legacy Title", "数量"])
            sheet.append(["PAWGO-19", "B0PAWGO019", self.payload["rows"][0]["legacy_title"], 2])
            sheet["A1"].fill = PatternFill("solid", fgColor="FFCC00")
            sheet["D2"] = "=1+1"
            workbook.create_sheet("说明")["A1"] = "必须保留"
            workbook.save(source)

            result = build_workbook(self.payload, source_file=str(source))

        self.assertEqual(result.sheetnames, ["原始数据", "说明"])
        self.assertEqual(result["说明"]["A1"].value, "必须保留")
        self.assertEqual(result["原始数据"]["D2"].value, "=1+1")
        self.assertEqual(result["原始数据"]["A1"].fill.fgColor.rgb, "00FFCC00")

    def test_appended_columns_are_highlighted_in_source_mode(self) -> None:
        sheet = build_workbook(self.payload, source_file=str(SOURCE_CSV))["标题结果"]
        headers = [cell.value for cell in sheet[1]]
        first_new = headers.index("新标题") + 1
        original = headers.index("SKU") + 1

        # 表头：原表列深蓝，新增列深绿
        self.assertEqual(sheet.cell(row=1, column=original).fill.fgColor.rgb, "001F4E79")
        self.assertEqual(sheet.cell(row=1, column=first_new).fill.fgColor.rgb, "002E7D32")

        # 数据区：原表单元格不着色，新增列浅绿底
        self.assertNotEqual(sheet.cell(row=2, column=original).fill.fgColor.rgb, "00E8F5E9")
        self.assertEqual(sheet.cell(row=2, column=first_new).fill.fgColor.rgb, "00E8F5E9")

    def test_legacy_title_column_is_dropped_when_source_already_has_it(self) -> None:
        import csv as _csv
        import tempfile as _tempfile

        rows = self.payload["rows"]
        with _tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "user_upload.csv"
            with open(source, "w", encoding="utf-8", newline="") as handle:
                writer = _csv.writer(handle)
                writer.writerow(["ASIN", "Original Title"])
                for row in rows:
                    writer.writerow([row.get("asin") or "", row.get("legacy_title") or ""])

            sheet = build_workbook(self.payload, source_file=str(source))["标题结果"]
            headers = [cell.value for cell in sheet[1]]

        # 用户表已自带原标题（列名不同但内容一致）→ 不重复追加
        self.assertNotIn("原标题", headers)
        self.assertIn("新标题", headers)
        self.assertEqual(headers[:2], ["ASIN", "Original Title"])

    def test_standalone_workbook_has_no_highlight(self) -> None:
        sheet = build_workbook(self.payload)["标题结果"]
        self.assertEqual(sheet.cell(row=1, column=1).fill.fgColor.rgb, "001F4E79")

    def test_workbook_saves_to_disk(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp) / "out.xlsx"
            build_workbook(self.payload).save(out)
            self.assertTrue(out.exists())
            self.assertGreater(out.stat().st_size, 4000)

    def test_payload_carries_xlsx_artifact_path(self) -> None:
        payload = _workbench(BATCH)
        self.assertIsNone(payload["artifacts"]["xlsx"])
        with tempfile.TemporaryDirectory() as tmp:
            out = str(Path(tmp) / "wb.xlsx")
            path = attach_xlsx_artifact(payload, out=out)
            self.assertEqual(Path(path).resolve(), Path(out).resolve())
            self.assertEqual(payload["artifacts"]["xlsx"], path)
            self.assertTrue(Path(out).exists())

    def test_strict_xlsx_failure_is_not_silently_swallowed(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            unsupported = Path(tmp) / "source.xls"
            unsupported.write_text("not an xlsx", encoding="utf-8")
            with self.assertRaisesRegex(RuntimeError, "XLSX 导出失败"):
                attach_xlsx_artifact(
                    _workbench(BATCH), source_file=str(unsupported), strict=True
                )


if __name__ == "__main__":
    unittest.main()
